"""Build the rulebook spreadsheet: a summary, then one tab per CSV in rulebook/.

The CSVs are the source of truth, and tests/test_rulebook.py checks them against the code.
The .xlsx is only a view for reading, filtering and sharing. It needs openpyxl, which the
engine itself does not:

    pip install openpyxl
    python tools/build_rulebook.py OpenPPC-rulebook.xlsx [--gold private/gold.csv]
"""
import argparse
import csv
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BOOK = Path(__file__).resolve().parent.parent / "rulebook"
TABS = (("Rules", "rules.csv"), ("Metrics", "metrics.csv"), ("Sources", "sources.csv"), ("Taxonomy", "taxonomy.csv"))
WIDTHS = {"rule": 60, "formula": 46, "calibration": 46, "notes": 44, "question": 46, "method": 50, "result": 60,
          "caveats": 60, "title": 50, "url": 48, "data_needed": 34, "gates": 40, "code_ref": 38, "threshold": 24,
          "conflicts_with": 22, "source_ids": 24, "name": 24, "sample": 34}
FILLS = {"live": "D9F2E3", "planned": "FFF1CC", "reference": "E8E8E8", "paid": "E9DDFB"}
HEADER = PatternFill("solid", fgColor="1F2937")
EVIDENCE = (("google", "Google's own guidance: Help Center, Skillshop, official guides"),
            ("expert", "Practitioner consensus, with a cited source"),
            ("openppc", "Our own rule, reasoned in code: read it and argue"),
            ("benchmark", "Public industry averages"),
            ("gold (paid)", "The calibration column: the same rule tuned on live accounts"))


def read(path):
    with open(path, newline="", encoding="utf-8") as f:
        return list(csv.reader(f))


def tab(wb, title, table):
    ws = wb.create_sheet(title)
    for row in table:
        ws.append(row)
    head = table[0]
    for cell in ws[1]:
        cell.font, cell.fill = Font(bold=True, color="FFFFFF"), HEADER
    for i, name in enumerate(head, 1):
        ws.column_dimensions[get_column_letter(i)].width = WIDTHS.get(name, 16)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for key in ("status", "tier"):
            if key in head and (fill := FILLS.get(row[head.index(key)].value)):
                row[head.index(key)].fill = PatternFill("solid", fgColor=fill)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions


def summary(ws, rules):
    body = [dict(zip(rules[0], r)) for r in rules[1:]]
    ws.title = "Summary"
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 22, 70
    ws.append(["OpenPPC rulebook"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append(["Every formula the tool uses, where it comes from, and what the paid tier calibrates. "
               "The CSVs in rulebook/ are the source; a test fails when they and the code disagree."])
    blocks = (("Rules by status and tier", Counter(f"{r['status']}, {r['tier']}" for r in body)),
              ("Rules by evidence", Counter(r["evidence"] for r in body)),
              ("Rules by area", Counter(r["area"] for r in body)))
    for title, counts in blocks:
        ws.append([])
        ws.append([title])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
        for key, n in counts.most_common():
            ws.append([key, n])
    ws.append([])
    ws.append(["Coverage by campaign type", "live", "planned", "reference"])
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
    cover = Counter((t.strip(), r["status"]) for r in body for t in r["campaign_type"].split(";"))
    for t in sorted({t for t, _ in cover}, key=lambda t: -sum(n for (k, _), n in cover.items() if k == t)):
        ws.append([t] + [cover[(t, s)] for s in ("live", "planned", "reference")])
    ws.append([])
    ws.append(["Evidence tiers"])
    ws.cell(ws.max_row, 1).font = Font(bold=True)
    for row in EVIDENCE:
        ws.append(list(row))


def main():
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("out")
    ap.add_argument("--gold", help="private calibration log (never commit it)")
    args = ap.parse_args()
    wb = Workbook()
    summary(wb.active, read(BOOK / "rules.csv"))
    for title, name in TABS:
        if (BOOK / name).exists():
            tab(wb, title, read(BOOK / name))
    if args.gold:
        tab(wb, "Gold (private)", read(args.gold))
    wb.save(args.out)
    print(f"wrote {args.out}")


if __name__ == "__main__":
    main()
